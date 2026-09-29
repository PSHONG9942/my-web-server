#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Self-contained PDF to PPTX Conversion Engine for FastAPI Backend Service
"""

import sys
import os
import io
import time
import base64
from pathlib import Path
from typing import Optional, List, Dict, Any, Callable

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    import pymupdf
except ImportError:
    try:
        import fitz as pymupdf
    except ImportError:
        pymupdf = None

try:
    import pptx
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
except ImportError:
    pptx = None


def parse_page_ranges(range_str: str, total_pages: int) -> List[int]:
    if not range_str or range_str.strip().lower() == "all":
        return list(range(total_pages))

    pages = set()
    parts = range_str.split(",")
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            sub = part.split("-")
            if len(sub) == 2:
                try:
                    start = int(sub[0].strip())
                    end = int(sub[1].strip())
                    for p in range(max(1, start), min(total_pages, end) + 1):
                        pages.add(p - 1)
                except ValueError:
                    pass
        else:
            try:
                p = int(part)
                if 1 <= p <= total_pages:
                    pages.add(p - 1)
            except ValueError:
                pass

    sorted_pages = sorted(list(pages))
    return sorted_pages if sorted_pages else list(range(total_pages))


def get_pdf_info(pdf_source) -> Dict[str, Any]:
    if pymupdf is None:
        raise RuntimeError("PyMuPDF is not installed. Please run: pip install pymupdf")

    if isinstance(pdf_source, (str, Path)):
        doc = pymupdf.open(str(pdf_source))
        file_size = os.path.getsize(str(pdf_source))
        file_name = Path(pdf_source).name
    elif isinstance(pdf_source, (bytes, bytearray)):
        doc = pymupdf.open(stream=pdf_source, filetype="pdf")
        file_size = len(pdf_source)
        file_name = "uploaded.pdf"
    else:
        raise ValueError("Invalid pdf_source type")

    total_pages = len(doc)
    if total_pages == 0:
        return {
            "file_name": file_name,
            "file_size": file_size,
            "total_pages": 0,
            "is_encrypted": doc.is_encrypted,
            "dimensions": [],
            "first_page_thumbnail": ""
        }

    first_page = doc[0]
    first_rect = first_page.rect
    is_rotated = first_page.rotation in (90, 270)
    w_pt = first_rect.height if is_rotated else first_rect.width
    h_pt = first_rect.width if is_rotated else first_rect.height
    aspect_ratio = round(w_pt / h_pt, 2) if h_pt > 0 else 1.78

    aspect_desc = "16:9 宽屏" if abs(aspect_ratio - 1.78) < 0.1 else (
        "4:3 标清" if abs(aspect_ratio - 1.33) < 0.1 else f"{w_pt:.0f} × {h_pt:.0f} pt"
    )

    thumb_pix = first_page.get_pixmap(dpi=90)
    thumb_bytes = thumb_pix.tobytes("jpeg", jpg_quality=80)
    thumb_b64 = "data:image/jpeg;base64," + base64.b64encode(thumb_bytes).decode("ascii")

    dimensions = []
    for i in range(min(5, total_pages)):
        p = doc[i]
        rot = p.rotation in (90, 270)
        p_w = p.rect.height if rot else p.rect.width
        p_h = p.rect.width if rot else p.rect.height
        dimensions.append({"page": i + 1, "width": round(p_w, 1), "height": round(p_h, 1)})

    title = doc.metadata.get("title", "") if doc.metadata else ""

    return {
        "file_name": file_name,
        "file_size": file_size,
        "total_pages": total_pages,
        "title": title or file_name,
        "is_encrypted": doc.is_encrypted,
        "aspect_ratio": aspect_ratio,
        "aspect_desc": aspect_desc,
        "slide_width_pt": round(w_pt, 1),
        "slide_height_pt": round(h_pt, 1),
        "dimensions": dimensions,
        "first_page_thumbnail": thumb_b64
    }


def convert_pdf_to_pptx(
    pdf_source,
    output_dest,
    mode: str = "presentation",
    dpi: int = 180,
    page_range: str = "all",
    progress_callback: Optional[Callable[[int, int, str, float], None]] = None
) -> Dict[str, Any]:
    if pymupdf is None:
        raise RuntimeError("PyMuPDF is required. Run: pip install pymupdf")
    if pptx is None:
        raise RuntimeError("python-pptx is required. Run: pip install python-pptx")

    t_start = time.time()

    if isinstance(pdf_source, (str, Path)):
        doc = pymupdf.open(str(pdf_source))
    elif isinstance(pdf_source, (bytes, bytearray)):
        doc = pymupdf.open(stream=pdf_source, filetype="pdf")
    else:
        raise ValueError("Invalid pdf_source type")

    total_doc_pages = len(doc)
    if total_doc_pages == 0:
        raise ValueError("The PDF document contains no pages.")

    if doc.is_encrypted:
        raise ValueError("The PDF document is password-protected or encrypted.")

    target_pages = parse_page_ranges(page_range, total_doc_pages)
    num_to_convert = len(target_pages)

    if progress_callback:
        progress_callback(0, num_to_convert, "初始化幻灯片结构...", 5.0)

    prs = pptx.Presentation()
    blank_layout = prs.slide_layouts[6]

    first_page = doc[target_pages[0]]
    is_rotated = first_page.rotation in (90, 270)
    base_w = first_page.rect.height if is_rotated else first_page.rect.width
    base_h = first_page.rect.width if is_rotated else first_page.rect.height

    prs.slide_width = Pt(base_w)
    prs.slide_height = Pt(base_h)

    for idx, page_idx in enumerate(target_pages):
        page = doc[page_idx]
        current_num = idx + 1
        pct = 5.0 + (idx / num_to_convert) * 85.0

        if progress_callback:
            progress_callback(current_num, num_to_convert, f"正在渲染第 {page_idx + 1} 页 ({current_num}/{num_to_convert})...", pct)

        p_rotated = page.rotation in (90, 270)
        p_w = page.rect.height if p_rotated else page.rect.width
        p_h = page.rect.width if p_rotated else page.rect.height

        slide = prs.slides.add_slide(blank_layout)

        if mode in ("presentation", "hybrid"):
            pix = page.get_pixmap(dpi=dpi)
            img_bytes = pix.tobytes("jpeg", jpg_quality=90)
            img_stream = io.BytesIO(img_bytes)
            slide.shapes.add_picture(img_stream, 0, 0, width=Pt(p_w), height=Pt(p_h))

        if mode in ("hybrid", "extracted"):
            try:
                page_dict = page.get_text("dict")
                blocks = page_dict.get("blocks", [])

                for b in blocks:
                    if b.get("type") == 0:
                        bbox = b.get("bbox", (0, 0, 0, 0))
                        bx0, by0, bx1, by1 = bbox
                        bw = max(10, bx1 - bx0)
                        bh = max(10, by1 - by0)

                        lines = b.get("lines", [])
                        block_text = ""
                        primary_size = 12
                        primary_color = (0, 0, 0)
                        is_bold = False

                        for l in lines:
                            for s in l.get("spans", []):
                                text_piece = s.get("text", "")
                                if text_piece.strip():
                                    block_text += text_piece + " "
                                    primary_size = s.get("size", primary_size)
                                    color_int = s.get("color", 0)
                                    primary_color = (
                                        (color_int >> 16) & 255,
                                        (color_int >> 8) & 255,
                                        color_int & 255
                                    )
                                    if s.get("flags", 0) & 2:
                                        is_bold = True
                            block_text += "\n"

                        block_text = block_text.strip()
                        if block_text:
                            tx_box = slide.shapes.add_textbox(Pt(bx0), Pt(by0), Pt(bw), Pt(bh))
                            tf = tx_box.text_frame
                            tf.word_wrap = True
                            tf.margin_left = Pt(2)
                            tf.margin_top = Pt(2)
                            tf.margin_right = Pt(2)
                            tf.margin_bottom = Pt(2)

                            p = tf.paragraphs[0]
                            p.text = block_text
                            p.font.size = Pt(min(60, max(8, primary_size)))
                            p.font.color.rgb = RGBColor(*primary_color)
                            p.font.bold = is_bold

                    elif b.get("type") == 1 and mode == "extracted":
                        img_bbox = b.get("bbox", (0, 0, 0, 0))
                        ix0, iy0, ix1, iy1 = img_bbox
                        iw = max(10, ix1 - ix0)
                        ih = max(10, iy1 - iy0)
                        img_data = b.get("image")
                        if img_data:
                            try:
                                slide.shapes.add_picture(
                                    io.BytesIO(img_data),
                                    Pt(ix0),
                                    Pt(iy0),
                                    width=Pt(iw),
                                    height=Pt(ih)
                                )
                            except Exception:
                                pass
            except Exception as e:
                print(f"[Warning] Text extraction error on page {page_idx + 1}: {e}")

    if progress_callback:
        progress_callback(num_to_convert, num_to_convert, "正在保存与打包 PPTX 演示文稿...", 95.0)

    if isinstance(output_dest, (str, Path)):
        prs.save(str(output_dest))
        out_size = os.path.getsize(str(output_dest))
        out_path_str = str(output_dest)
    else:
        prs.save(output_dest)
        output_dest.seek(0)
        out_size = len(output_dest.getvalue())
        out_path_str = "memory_buffer"

    t_duration = round(time.time() - t_start, 2)

    if progress_callback:
        progress_callback(num_to_convert, num_to_convert, "转换完成！", 100.0)

    return {
        "success": True,
        "pages_converted": num_to_convert,
        "total_doc_pages": total_doc_pages,
        "duration_seconds": t_duration,
        "output_size_bytes": out_size,
        "output_path": out_path_str,
        "mode": mode,
        "dpi": dpi
    }
